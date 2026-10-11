"""User-supplied sources: CSV/XLSX lists, website URL lists and directory pages.

These are the legitimate path for sources that do not permit automated
collection (e.g. member directories behind logins, permission-gated datasets):
the user obtains the data under that source's terms and imports it here.
"""

from __future__ import annotations

import csv
import hashlib
import io
from collections.abc import Iterable
from urllib.parse import urljoin, urlsplit

from bs4 import BeautifulSoup

from atlas.models.company import CandidateCompany, CandidateLocation
from atlas.models.enums import ProviderAvailability
from atlas.models.search import SearchCriteria
from atlas.providers.base import (
    DiscoveryProvider,
    ProviderContext,
    ProviderError,
    ProviderInfo,
    ProviderResult,
)
from atlas.resolve.normalize import is_company_domain, normalize_url, registrable_domain
from atlas.timeutil import now_iso

# Accepted header names (lower-case) for each canonical column.
COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "name": ("name", "company", "company name", "company_name", "organisation", "organization", "trade name",
             "trade_name", "business name", "license name", "licence name"),
    "website": ("website", "url", "web", "site", "homepage", "web site", "domain"),
    "email": ("email", "e-mail", "email address", "contact email"),
    "phone": ("phone", "telephone", "tel", "phone number", "mobile", "contact number"),
    "address": ("address", "location", "full address", "street"),
    "city": ("city", "emirate", "town"),
    "industry": ("industry", "sector", "activity", "activities", "business activity", "category"),
    "notes": ("notes", "note", "comments", "description"),
}


def map_columns(headers: Iterable[str]) -> dict[str, str]:
    """Map canonical column → actual header (case-insensitive)."""
    lookup = {h.strip().lower(): h for h in headers if h}
    out: dict[str, str] = {}
    for canon, aliases in COLUMN_ALIASES.items():
        for a in aliases:
            if a in lookup:
                out[canon] = lookup[a]
                break
    return out


def rows_to_candidates(rows: list[dict[str, str]], *, source_label: str, mapping: dict[str, str] | None = None
                       ) -> tuple[list[CandidateCompany], list[str]]:
    if not rows:
        return [], ["file contains no rows"]
    mapping = mapping or map_columns(rows[0].keys())
    if "name" not in mapping and "website" not in mapping:
        raise ProviderError("The file needs at least a company name or website column. "
                            f"Recognised headers: {sorted(COLUMN_ALIASES['name'] + COLUMN_ALIASES['website'])}")
    out: list[CandidateCompany] = []
    problems: list[str] = []
    retrieved = now_iso()
    for i, row in enumerate(rows, start=2):  # row 1 is the header
        def val(col: str) -> str | None:
            h = mapping.get(col)
            v = (row.get(h) or "").strip() if h else ""
            return v or None

        website = normalize_url(val("website"))
        name = val("name")
        if not name and website:
            name = registrable_domain(website) or website
        if not name:
            problems.append(f"row {i}: no name or website — skipped")
            continue
        ext = hashlib.sha1(f"{source_label}|{name.lower()}|{website or ''}".encode()).hexdigest()[:16]
        out.append(CandidateCompany(
            name=name, provider="manual_csv", external_id=ext, source_url=None, website=website,
            email=val("email"), phone=val("phone"), description=val("notes"),
            location=CandidateLocation(address=val("address"), city=val("city"), country_code="AE")
            if (val("address") or val("city")) else None,
            industry_hints=[val("industry")] if val("industry") else [],
            tags={"import_source": source_label, "row": i, "industry_text": val("industry")},
            retrieved_at=retrieved, relevance_hint=0.5,
        ))
    return out, problems


def parse_tabular(data: bytes, filename: str) -> list[dict[str, str]]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        headers = [str(h).strip() if h is not None else "" for h in next(it, [])]
        return [{headers[i]: ("" if v is None else str(v)) for i, v in enumerate(r) if i < len(headers)} for r in it]
    text = data.decode("utf-8-sig", errors="replace")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return list(csv.DictReader(io.StringIO(text), dialect=dialect))


def urls_to_candidates(text: str) -> tuple[list[CandidateCompany], list[str]]:
    out, problems, seen = [], [], set()
    retrieved = now_iso()
    for line in text.splitlines():
        raw = line.strip().strip(",")
        if not raw or raw.startswith("#"):
            continue
        url = normalize_url(raw)
        dom = registrable_domain(url) if url else None
        if not url or not dom:
            problems.append(f"not a valid website URL: {raw[:80]}")
            continue
        if not is_company_domain(dom):
            problems.append(f"{dom} is a social/directory host, not a company website — skipped")
            continue
        if dom in seen:
            continue
        seen.add(dom)
        out.append(CandidateCompany(name=dom, provider="url_list", external_id=dom, website=url, source_url=url,
                                    retrieved_at=retrieved, relevance_hint=0.5,
                                    tags={"name_from_domain": True}))
    return out, problems


class ManualCsvProvider(DiscoveryProvider):
    info = ProviderInfo(
        name="manual_csv",
        display_name="Manual import (CSV / XLSX)",
        availability=ProviderAvailability.MANUAL_ONLY,
        description="Import a company list you obtained legitimately (e.g. a free-zone directory export, an "
                    "exhibitor list, a Dubai Pulse dataset you have permission to use).",
        coverage="Whatever you import.",
        source_url="",
        limitations="Imported rows are treated as unverified until their websites are researched.",
        enabled_by_default=True,
        runs_in_search=False,
    )

    async def discover(self, criteria: SearchCriteria, ctx: ProviderContext) -> ProviderResult:
        raise ProviderError("Manual import runs from the Search page, not as part of a search query")


class UrlListProvider(DiscoveryProvider):
    info = ProviderInfo(
        name="url_list",
        display_name="Company website list",
        availability=ProviderAvailability.MANUAL_ONLY,
        description="Paste company website URLs you already know; ATLAS researches each one.",
        coverage="Whatever you paste.",
        source_url="",
        limitations="Company names are taken from the website during research.",
        enabled_by_default=True,
        runs_in_search=False,
    )

    async def discover(self, criteria: SearchCriteria, ctx: ProviderContext) -> ProviderResult:
        raise ProviderError("URL lists are imported from the Search page")


async def directory_page_candidates(url: str, ctx: ProviderContext, *, max_links: int = 300
                                    ) -> tuple[list[CandidateCompany], list[str]]:
    """Extract outbound company-website links from a public directory page.

    robots.txt is honoured; the user must confirm the directory's terms allow
    this use before calling it (enforced in the UI).
    """
    page_url = normalize_url(url)
    if not page_url:
        raise ProviderError("invalid directory URL")
    res = await ctx.fetcher.get(page_url)
    if res.robots_blocked:
        raise ProviderError("robots.txt disallows fetching this directory page — use manual CSV import instead")
    if not res.ok:
        raise ProviderError(f"could not fetch directory page: {res.error or res.status}")
    soup = BeautifulSoup(res.text(), "lxml")
    page_dom = registrable_domain(res.final_url)
    out: list[CandidateCompany] = []
    seen: set[str] = set()
    notes: list[str] = []
    retrieved = now_iso()
    for a in soup.find_all("a", href=True):
        href = urljoin(res.final_url, a["href"])
        if not href.startswith(("http://", "https://")):
            continue
        dom = registrable_domain(href)
        if not dom or dom == page_dom or dom in seen or not is_company_domain(dom):
            continue
        text = " ".join(a.get_text(" ", strip=True).split())[:120]
        seen.add(dom)
        root = f"{urlsplit(href).scheme}://{urlsplit(href).netloc}/"
        out.append(CandidateCompany(
            name=text if 2 < len(text) < 100 and not text.lower().startswith(("http", "www", "visit")) else dom,
            provider="directory_page", external_id=dom, website=root, source_url=page_url,
            retrieved_at=retrieved, relevance_hint=0.4, tags={"directory": page_url, "anchor_text": text}))
        if len(out) >= max_links:
            notes.append(f"stopped at {max_links} links")
            break
    notes.append(f"{len(out)} distinct external company websites found on the page")
    return out, notes


class DirectoryPageProvider(DiscoveryProvider):
    info = ProviderInfo(
        name="directory_page",
        display_name="Public directory page (user-supplied URL)",
        availability=ProviderAvailability.MANUAL_ONLY,
        description="Extract company website links from a public list page you provide (e.g. an exhibitor list "
                    "or association member page) after you confirm its terms allow it.",
        coverage="One page per import.",
        source_url="",
        limitations="robots.txt is enforced. Pages behind logins or JavaScript-only listings are not supported.",
        enabled_by_default=True,
        runs_in_search=False,
    )

    async def discover(self, criteria: SearchCriteria, ctx: ProviderContext) -> ProviderResult:
        raise ProviderError("Directory pages are imported from the Search page")
